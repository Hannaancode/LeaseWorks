"""Bounded document reading with stable, human-readable source locations."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from docx import Document
from pypdf import PdfReader

MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_TEXT = 80000


def read_document(filename: str, raw: bytes) -> list[dict]:
    if not raw or len(raw) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document must be between 1 byte and 8 MiB")
    suffix = Path(filename).suffix.lower()
    segments = []

    def add(text, location):
        for line in text.splitlines():
            if line.strip():
                segments.append({"id": f"s{len(segments) + 1}", "text": line.strip(), "location": location})

    if suffix == ".txt":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("TXT must be UTF-8 encoded") from exc
        for i, line in enumerate(text.splitlines(), 1):
            add(line, f"Line {i}")
    elif suffix == ".pdf":
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 30:
            raise ValueError("Use an unencrypted PDF of at most 30 pages")
        for i, page in enumerate(reader.pages, 1):
            add(page.extract_text() or "", f"Page {i}")
    elif suffix == ".docx":
        try:
            with ZipFile(BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError("Expanded DOCX exceeds 20 MiB")
            doc = Document(BytesIO(raw))
        except BadZipFile as exc:
            raise ValueError("Invalid DOCX") from exc
        for i, p in enumerate(doc.paragraphs, 1):
            add(p.text, f"Paragraph {i}")
        for i, table in enumerate(doc.tables, 1):
            for j, row in enumerate(table.rows, 1):
                add(" | ".join(c.text for c in row.cells), f"Table {i}, row {j}")
    else:
        raise ValueError("Supported documents: TXT, text-based PDF and DOCX")
    if not segments:
        raise ValueError("No readable text. Scanned PDFs need OCR, which this prototype does not provide.")
    if sum(len(s["text"]) for s in segments) > MAX_TEXT:
        raise ValueError("Extracted text exceeds 80,000 characters")
    return segments
